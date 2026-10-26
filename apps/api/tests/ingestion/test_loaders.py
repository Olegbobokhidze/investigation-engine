from io import BytesIO
from pathlib import Path

import docx
import pytest

from investigation_engine.core.errors import (
    EmptyDocumentError,
    UnreadableFileError,
    UnsupportedFormatError,
)
from investigation_engine.documents.models import FileType
from investigation_engine.ingestion.loaders import load_document

FIXTURES = Path(__file__).parent.parent / "fixtures"


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.core_properties.title = "checkout-api runbook"
    document.core_properties.author = "Platform team"
    document.add_heading("Timeouts", level=1)
    document.add_paragraph("Fraud scoring times out after 3 seconds.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Retries"
    table.rows[0].cells[1].text = "1"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_load_pdf_extracts_pages_and_metadata() -> None:
    content = (FIXTURES / "incident-report.pdf").read_bytes()

    loaded = load_document(content, "incident-report.pdf")

    assert loaded.file_type == FileType.PDF
    assert "first alerts fired at 09:12" in loaded.text
    assert "\n\nResolution" in loaded.text  # pages stay separate paragraphs
    assert loaded.metadata["page_count"] == 2
    assert loaded.metadata["title"] == "INC-0412 Payment latency"


def test_load_docx_extracts_paragraphs_tables_and_metadata() -> None:
    loaded = load_document(_docx_bytes(), "runbook.docx")

    assert loaded.file_type == FileType.DOCX
    assert "Fraud scoring times out after 3 seconds." in loaded.text
    assert "Retries | 1" in loaded.text
    assert loaded.metadata["title"] == "checkout-api runbook"
    assert loaded.metadata["author"] == "Platform team"


def test_load_markdown_keeps_headings() -> None:
    content = b"# Change log\n\n## 14 March\n\nCHG-2291 deployed at 09:40."

    loaded = load_document(content, "CHANGELOG.MD")

    assert loaded.file_type == FileType.MARKDOWN
    assert loaded.text.startswith("# Change log\n\n## 14 March")


def test_load_text_strips_byte_order_mark_and_counts_words() -> None:
    content = "﻿Ledger lag peaked at 41 minutes.".encode()

    loaded = load_document(content, "notes.txt")

    assert loaded.file_type == FileType.TEXT
    assert loaded.text == "Ledger lag peaked at 41 minutes."
    assert loaded.metadata["word_count"] == 6


def test_load_document_rejects_unknown_extension() -> None:
    with pytest.raises(UnsupportedFormatError):
        load_document(b"<html></html>", "page.html")


def test_load_document_rejects_file_without_text() -> None:
    with pytest.raises(EmptyDocumentError):
        load_document(b"  \n\n \t ", "blank.txt")


@pytest.mark.parametrize(
    ("content", "file_name"),
    [
        (b"not a pdf", "broken.pdf"),
        (b"not a zip archive", "broken.docx"),
        (b"\xff\xfe\xfa invalid utf-8", "binary.txt"),
    ],
)
def test_load_document_rejects_unreadable_file(content: bytes, file_name: str) -> None:
    with pytest.raises(UnreadableFileError):
        load_document(content, file_name)
