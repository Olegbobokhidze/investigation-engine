"""Turns an uploaded file into plain text plus the metadata the format provides.

Each format is one `load_*` function registered in `LOADERS` by file extension; supporting
a new format means writing one function and adding its extension there.
"""

import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import docx
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from investigation_engine.core.errors import (
    EmptyDocumentError,
    UnreadableFileError,
    UnsupportedFormatError,
)
from investigation_engine.documents.models import FileType
from investigation_engine.ingestion.normalization import normalize_text

type Metadata = dict[str, str | int]


@dataclass(frozen=True)
class LoadedDocument:
    text: str
    file_type: FileType
    metadata: Metadata = field(default_factory=dict)


type Loader = Callable[[bytes], LoadedDocument]


def _decode_text(content: bytes) -> str:
    try:
        # utf-8-sig also strips the byte-order mark some editors write.
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise UnreadableFileError("The file isn't valid UTF-8 text.") from error


def load_pdf(content: bytes) -> LoadedDocument:
    try:
        reader = PdfReader(BytesIO(content))
        pages = [page.extract_text() for page in reader.pages]
    except PdfReadError as error:
        raise UnreadableFileError("The PDF is damaged or encrypted.") from error

    metadata: Metadata = {"page_count": len(pages)}
    if reader.metadata is not None:
        if reader.metadata.title:
            metadata["title"] = reader.metadata.title
        if reader.metadata.author:
            metadata["author"] = reader.metadata.author
    # A blank line between pages keeps them as separate paragraphs for chunking.
    return LoadedDocument(text="\n\n".join(pages), file_type=FileType.PDF, metadata=metadata)


def load_docx(content: bytes) -> LoadedDocument:
    try:
        document = docx.Document(BytesIO(content))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError) as error:
        raise UnreadableFileError("The DOCX file is damaged.") from error

    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    # Tables hold real content in reports and runbooks; keep each row as one line.
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))

    metadata: Metadata = {}
    properties = document.core_properties
    if properties.title:
        metadata["title"] = properties.title
    if properties.author:
        metadata["author"] = properties.author
    return LoadedDocument(text="\n\n".join(paragraphs), file_type=FileType.DOCX, metadata=metadata)


def load_markdown(content: bytes) -> LoadedDocument:
    # Markup stays in: headings mark section boundaries for the structure-aware chunker.
    return LoadedDocument(text=_decode_text(content), file_type=FileType.MARKDOWN)


def load_text(content: bytes) -> LoadedDocument:
    return LoadedDocument(text=_decode_text(content), file_type=FileType.TEXT)


LOADERS: dict[str, Loader] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".md": load_markdown,
    ".markdown": load_markdown,
    ".txt": load_text,
}


def load_document(content: bytes, file_name: str) -> LoadedDocument:
    """Picks the loader by extension, then normalizes the text and adds common metadata.

    Raises `UnsupportedFormatError`, `UnreadableFileError` or `EmptyDocumentError`.
    """
    extension = Path(file_name).suffix.lower()
    loader = LOADERS.get(extension)
    if loader is None:
        supported = ", ".join(sorted(LOADERS))
        raise UnsupportedFormatError(f"Unsupported file type “{extension}”. Use {supported}.")

    loaded = loader(content)
    text = normalize_text(loaded.text)
    if not text:
        raise EmptyDocumentError("The file contains no readable text.")

    metadata: Metadata = {
        **loaded.metadata,
        "character_count": len(text),
        "word_count": len(text.split()),
    }
    return LoadedDocument(text=text, file_type=loaded.file_type, metadata=metadata)
